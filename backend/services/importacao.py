import csv
import io
import logging
from datetime import datetime
from typing import Optional

import openpyxl
from sqlalchemy.orm import Session

import models

logger = logging.getLogger(__name__)


def _calcular_stats(dezenas: list[int], anterior: Optional[list[int]]) -> dict:
    total_pares = sum(1 for d in dezenas if d % 2 == 0)
    total_impares = 15 - total_pares
    repetidas = len(set(dezenas) & set(anterior)) if anterior else 0
    return {
        "total_pares": total_pares,
        "total_impares": total_impares,
        "repetidas_anterior": repetidas,
    }


def _normalizar_header(h: str) -> str:
    return (
        h.lower()
        .strip()
        .replace(" ", "")
        .replace("ª", "")
        .replace("º", "")
        .replace(".", "")
        .replace("_", "")
    )


def _detectar_colunas(headers: list[str]) -> tuple[Optional[list[str]], Optional[str]]:
    """Retorna (bola_cols, dezenas_col). Um dos dois será None."""
    norm = [_normalizar_header(h) for h in headers]
    norm_to_original = dict(zip(norm, headers))

    # Prefixos conhecidos: bola1..bola15, dezena1..dezena15
    for prefix in ("bola", "dezena", "d"):
        cols = []
        for i in range(1, 16):
            key = f"{prefix}{i}"
            if key in norm:
                cols.append(norm_to_original[key])
        if len(cols) == 15:
            return cols, None

    # Sufixo invertido: 1dezena, 2bola, 1bola...
    for suffix in ("dezena", "bola"):
        cols = []
        for i in range(1, 16):
            key = f"{i}{suffix}"
            if key in norm:
                cols.append(norm_to_original[key])
        if len(cols) == 15:
            return cols, None

    # Coluna única com dezenas separadas
    for pattern in ("dezenas", "númerossorteados", "numerossorteados", "bolas", "numeros"):
        if pattern in norm:
            return None, norm_to_original[pattern]

    return None, None


def _encontrar_coluna(headers: list[str], candidatos: list[str]) -> str:
    norm = [_normalizar_header(h) for h in headers]
    norm_to_original = dict(zip(norm, headers))
    for c in candidatos:
        if c in norm:
            return norm_to_original[c]
    return headers[0]


def _coluna_concurso(headers: list[str]) -> str:
    return _encontrar_coluna(
        headers,
        ["concurso", "numeroconcurso", "númerodoconcurso", "nconcurso", "noconcurso"],
    )


def _coluna_data(headers: list[str]) -> str:
    return _encontrar_coluna(
        headers,
        ["datasorteio", "datadosorteio", "data", "datasorteados"],
    )


def _parse_data(raw) -> datetime:
    if isinstance(raw, datetime):
        return raw
    s = str(raw).strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d/%m/%y"):
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            pass
    raise ValueError(f"Formato de data não reconhecido: {s!r}")


def _parse_dezenas(row: dict, bola_cols: Optional[list[str]], dezenas_col: Optional[str]) -> list[int]:
    if bola_cols:
        return sorted(int(row[c]) for c in bola_cols)
    raw = str(row[dezenas_col]).strip()
    parts = raw.replace(",", " ").replace(";", " ").split()
    return sorted(int(p) for p in parts)


def _detectar_delimitador(amostra: str) -> str:
    try:
        dialect = csv.Sniffer().sniff(amostra, delimiters=",;|\t")
        return dialect.delimiter
    except csv.Error:
        return ","


def _linha_e_cabecalho(row: tuple) -> bool:
    """Identifica a linha real de cabecalhos, mesmo com texto explicativo antes
    (comum em planilhas baixadas de sites como asloterias.com.br)."""
    norm_cells = {_normalizar_header(str(c)) for c in row if c is not None}
    return "concurso" in norm_cells


def _ler_xlsx(conteudo: bytes) -> tuple[list[str], list[dict]]:
    wb = openpyxl.load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
    sheet_names_upper = [s.upper() for s in wb.sheetnames]
    if "LOTOFÁCIL" in sheet_names_upper or "LOTOFACIL" in sheet_names_upper:
        idx = next(
            i for i, s in enumerate(sheet_names_upper)
            if s in ("LOTOFÁCIL", "LOTOFACIL")
        )
        ws = wb[wb.sheetnames[idx]]
    else:
        ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        return [], []

    header_idx = 0
    for i, row in enumerate(rows[:30]):
        if any(c is not None for c in row) and _linha_e_cabecalho(row):
            header_idx = i
            break

    headers = [str(c) if c is not None else "" for c in rows[header_idx]]
    data = []
    for row in rows[header_idx + 1:]:
        if all(c is None for c in row):
            continue
        data.append(dict(zip(headers, row)))
    return headers, data


def _ler_csv(conteudo: bytes) -> tuple[list[str], list[dict]]:
    text = conteudo.decode("utf-8-sig")
    delimitador = _detectar_delimitador(text[:4096])
    reader = csv.DictReader(io.StringIO(text), delimiter=delimitador)
    headers = list(reader.fieldnames or [])
    data = [row for row in reader if any(v.strip() for v in row.values())]
    return headers, data


def importar_historico(conteudo: bytes, filename: str, db: Session) -> dict:
    if filename.lower().endswith(".csv"):
        headers, rows = _ler_csv(conteudo)
    else:
        headers, rows = _ler_xlsx(conteudo)

    total_lidas = len(rows)
    logger.info("Arquivo lido: %d linhas, cabeçalhos: %s", total_lidas, headers)

    if not rows:
        return {"total_lidas": 0, "total_inseridos": 0, "total_ignorados": 0}

    bola_cols, dezenas_col = _detectar_colunas(headers)
    if bola_cols is None and dezenas_col is None:
        raise ValueError(
            f"Não foi possível detectar colunas de dezenas. "
            f"Cabeçalhos encontrados: {headers}"
        )

    col_concurso = _coluna_concurso(headers)
    col_data = _coluna_data(headers)

    # Parse e valida cada linha
    registros: list[dict] = []
    for row in rows:
        try:
            numero = int(row[col_concurso])
            data = _parse_data(row[col_data])
            dezenas = _parse_dezenas(row, bola_cols, dezenas_col)
            if len(dezenas) != 15 or not all(1 <= d <= 25 for d in dezenas):
                continue
            registros.append({"numero_concurso": numero, "data_sorteio": data, "dezenas": dezenas})
        except (ValueError, TypeError, KeyError):
            continue

    # Ordena pelo número do concurso para calcular repetidas_anterior corretamente
    registros.sort(key=lambda r: r["numero_concurso"])

    # Uma única query para obter todos os concursos existentes
    existentes: set[int] = {
        row[0] for row in db.query(models.Sorteio.numero_concurso).all()
    }

    # Deduplica dentro do próprio arquivo (mantém primeira ocorrência)
    vistos: set[int] = set()
    registros_unicos: list[dict] = []
    for r in registros:
        if r["numero_concurso"] not in vistos:
            vistos.add(r["numero_concurso"])
            registros_unicos.append(r)

    novos = [r for r in registros_unicos if r["numero_concurso"] not in existentes]
    total_ignorados = total_lidas - len(novos)

    if not novos:
        return {
            "total_lidas": total_lidas,
            "total_inseridos": 0,
            "total_ignorados": total_lidas,
        }

    # Carrega dezenas do banco (para calcular repetidas_anterior do primeiro lote)
    dezenas_por_concurso: dict[int, list[int]] = {
        num: dez
        for num, dez in db.query(models.Sorteio.numero_concurso, models.Sorteio.dezenas).all()
    }

    # Monta objetos para inserção em lote
    inserir = []
    for r in novos:
        anterior = dezenas_por_concurso.get(r["numero_concurso"] - 1)
        stats = _calcular_stats(r["dezenas"], anterior)
        inserir.append(
            models.Sorteio(
                numero_concurso=r["numero_concurso"],
                data_sorteio=r["data_sorteio"],
                dezenas=r["dezenas"],
                **stats,
            )
        )
        # Disponibiliza para os próximos da fila calcularem repetidas_anterior
        dezenas_por_concurso[r["numero_concurso"]] = r["dezenas"]

    db.bulk_save_objects(inserir)
    db.commit()

    logger.info("Importação concluída: %d inseridos, %d ignorados.", len(inserir), total_ignorados)
    return {
        "total_lidas": total_lidas,
        "total_inseridos": len(inserir),
        "total_ignorados": total_ignorados,
    }
